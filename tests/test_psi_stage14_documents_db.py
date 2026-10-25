"""PSI Stage14 fixture A: real PDF + DOCX -> persisted segments -> grounded
Evidence -> human approval -> immutable Context -> budgeted mock RTB -> SRE.

No paid provider calls. Both document formats are generated locally and
extracted with the same production parser inside a mocked process boundary.
"""
from __future__ import annotations

import base64
import io
import json
import os
import unittest
from unittest.mock import patch

from test_rtb_context_stage12_flow_db import ApprovedContextRTBFlow
from intelligence.project_sources import file_library, document_jobs, evidence_jobs, context_service
from intelligence.project_sources.document_parser import extract
from intelligence.taxonomy_builder.cost_repository import create_quote, authorize_budget
from intelligence.taxonomy_builder.run_repository import (
    start_run, advance_run, submit_fact_review, publish_reviewed_topics, get_run)
from test_taxonomy_builder_run_db import reply


@unittest.skipUnless(os.getenv("CM_TEST_PG_DSN"), "PostgreSQL17 required")
class ProjectDocumentsToTaxonomy(ApprovedContextRTBFlow):
    def test_pdf_docx_grounded_context_to_taxonomy(self):
        import fitz
        from docx import Document
        from psycopg import connect
        from pathlib import Path

        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "People struggle to track their daily meals.")
        pdf_bytes = pdf.tobytes()
        pdf.close()

        document = Document()
        document.add_heading("Audience Research", 1)
        document.add_paragraph("Users want a faster picture-based food log.")
        docx_output = io.BytesIO()
        document.save(docx_output)

        blob_sources = [
            ("evidence.pdf", pdf_bytes, "People struggle to track their daily meals."),
            ("evidence.docx", docx_output.getvalue(), "Users want a faster picture-based food log.")
        ]
        approved_docs = []
        with patch.object(file_library, "_project", return_value=self.project), patch.object(
                file_library, "_connect", side_effect=lambda: connect(self.dsn)), patch(
                "projects.storage_root", return_value=Path(self.tmp.name)), patch.object(
                document_jobs, "isolated_extract",
                side_effect=lambda blob,name,cfg: extract(blob,name,config=cfg)):
            for filename, blob, fragment in blob_sources:
                fid = file_library.add_file(self.pid, {
                    "name":filename,"data_base64":base64.b64encode(blob).decode(),
                    "uploaded_by":"stage14-offline-fixture"})["file_id"]
                task = document_jobs.enqueue(self.pid,fid,{"config":{"ocr":False}})
                self.assertEqual(task["state"],"queued")
                self.assertEqual(document_jobs.run_one()["state"],"completed")
                result = document_jobs.get_result(self.pid,fid,task["job_id"])
                self.assertTrue(any(fragment in s["text"] for s in result["result"]["segments"]))
                ev_task = evidence_jobs.enqueue(self.pid,fid,{})
                self.assertEqual(ev_task["state"],"queued")
                self.assertEqual(evidence_jobs.run_one()["state"],"completed")
                entries = evidence_jobs.list_evidence(self.pid,sid=fid,limit=100)["items"]
                candidates = [e for e in entries if fragment in e["exact_quote"]]
                self.assertTrue(candidates,(filename,entries))
                for item in candidates:
                    self.assertEqual(item["trust_label"],"extracted_document_unverified")
                    self.assertTrue(item["citation"].startswith("project-source://"))
                    evidence_jobs.review(self.pid,item["evidence_id"],{"decision":"approved"})
                    approved_docs.append(item["evidence_id"])

        initial=context_service.preview(self.pid)
        self.assertEqual(initial["context"]["coverage"]["sources_incomplete"],0)
        snap=context_service.create_snapshot(
            self.pid,{"expected_context_sha256":initial["context_sha256"]})
        context_service.decide(self.pid,snap["context_id"],{
            "decision":"approved","expected_context_sha256":initial["context_sha256"],
            "acknowledge_partial":False,"excluded_source_ids":[]})
        quote=create_quote(self.pid,self.pipeline["pipeline_id"],
            context_id=snap["context_id"],token_budget=5500,force_new=True)["quote"]
        chosen=quote["details"]["context_binding"]["selected_evidence_ids"]
        for eid in approved_docs:self.assertIn(eid,chosen)
        consent=authorize_budget(self.pid,quote["quote_id"],"0.30")["consent"]
        rid=start_run(self.pid,quote["quote_id"],consent["consent_id"],
                      "stage14-pdf-docx-context-e2e-v1")["run_id"]
        self.assertEqual(get_run(self.pid,rid)["input_mode"],"approved_project_context_v1")
        transmitted=[]
        def local_sender(payload):
            role=payload["response_format"]["json_schema"]["name"].removeprefix("tb_")
            request=json.loads(payload["messages"][1]["content"])
            transmitted.append((role,request))
            if role!="fact_extraction":
                return self.sender(payload)
            rows=request["context_evidence"]
            profile=next(e for e in rows if "Считает калории по фото" in e["text"])
            problem=next(e for e in rows if "Трудно отслеживать калории" in e["text"])
            document_row=next(e for e in rows if
                "People struggle to track their daily meals." in e["text"])
            body={"facts":[
                {"kind":"function","text":"Photo-based calorie estimation",
                 "source_evidence_id":profile["evidence_id"],"source_quote":"Считает калории по фото"},
                {"kind":"problem","text":"Tracking meals is inconvenient",
                 "source_evidence_id":problem["evidence_id"],"source_quote":"Трудно отслеживать калории"},
                {"kind":"problem","text":"Some people struggle to track meals",
                 "source_evidence_id":document_row["evidence_id"],
                 "source_quote":"People struggle to track their daily meals."}]}
            out=reply(role,request)
            out["model"]=payload["model"]
            out["choices"][0]["message"]["content"]=json.dumps(body,ensure_ascii=False)
            return out
        self.assertEqual(advance_run(self.pid,rid,provider_call=local_sender)["status"],"fact_review")
        proposed=get_run(self.pid,rid)["fact_proposals"]
        self.assertEqual(len(proposed),3)
        cited=next(f for f in proposed if f.get("source_evidence_id") in approved_docs)
        self.assertEqual(cited["kind"],"problem")
        self.assertEqual(cited["source_type"],"research_document")
        self.assertEqual(cited["trust_label"],"extracted_document_unverified")
        submit_fact_review(self.pid,rid,[{"id":x["id"],"action":"approve"} for x in proposed])
        self.assertEqual(advance_run(self.pid,rid,provider_call=local_sender)["status"],"topic_review")
        topics=get_run(self.pid,rid)["topic_proposals"]
        published=publish_reviewed_topics(self.pid,rid,[
            {"id":x["id"],"action":"approve"} for x in topics])
        self.assertEqual(published["status"],"approved")
        self.assertEqual(get_run(self.pid,rid)["publication"]["version"],1)
        self.assertTrue(any(x.get("source_type")=="research_document"
                            for x in transmitted[0][1]["context_evidence"]))
        self.assertEqual([x[0] for x in transmitted],
                         ["fact_extraction","taxonomy_generation"])


if __name__=="__main__":
    unittest.main()
